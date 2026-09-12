from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('documentation')
OUT.mkdir(exist_ok=True)
NAVY = '17365D'; PALE = 'EAF2F8'; GRAY = 'F3F5F7'; BORDER = 'D9D9D9'

def shade(cell, fill):
    tc = cell._tc
    pr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); pr.append(shd)

def borders(table):
    pr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),BORDER); b.append(e)
    pr.append(b)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,NAVY); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.bold=True; r.font.size=Pt(9)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],GRAY)
            for p in cells[i].paragraphs:
                for r in p.runs:r.font.size=Pt(8.5)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t

def setup(doc, title, subtitle):
    sec=doc.sections[0]; sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal'].font.color.rgb=RGBColor(25,25,25)
    styles['Normal'].paragraph_format.space_after=Pt(6); styles['Normal'].paragraph_format.line_spacing=1.08
    for name,size in [('Title',26),('Heading 1',17),('Heading 2',12)]:
        styles[name].font.name='Aptos Display'; styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor(0,0,0); styles[name].font.bold=True
    title_pr = styles['Title']._element.get_or_add_pPr()
    for border in title_pr.findall(qn('w:pBdr')):
        title_pr.remove(border)
    p=doc.add_paragraph(style='Title'); p.add_run(title)
    p=doc.add_paragraph(); p.add_run(subtitle).bold=True
    doc.add_paragraph('Prepared by David Gaisey-Otoo | AWS Career Launchpad Pro | 12 September 2026')

def bullets(doc, items):
    for item in items: doc.add_paragraph(item, style='List Bullet')

def numbered(doc, items):
    for item in items: doc.add_paragraph(item, style='List Number')

def build_main():
    d=Document(); setup(d,'AWS Identity and Account Security Assessment','A to Z portfolio delivery dossier and practical runbook')
    d.add_heading('Purpose and conclusion',1)
    d.add_paragraph('This self-directed portfolio engagement assesses an existing AWS training account through read-only console and CloudShell checks. It creates no AWS resources, does not change permissions, and converts verified results into a redacted security report and portfolio case study. Generated plans are not evidence; each completed claim requires an actual result and screenshot reference.')
    table(d,['Field','Approved value'],[
        ['Project owner','David Gaisey-Otoo'],['Engagement type','Self-directed AWS portfolio assessment'],['Environment','Existing AWS training account'],['Execution mode','Read-only AWS assessment'],['Resource creation','Prohibited'],['Target duration','One working week'],['Success criterion','Evidence-backed findings, no secrets, and proof that no resources were created']],[2.0,4.9])
    d.add_heading('Proposal and scope agreement',1)
    d.add_paragraph('Proposal. David Gaisey-Otoo will perform a read-only identity, resource-inventory, audit-history, and cost-governance assessment. The work will produce an architecture specification, command log, test record, findings register, external-review package, handover record, and portfolio-safe case study.')
    d.add_heading('Included work',2); bullets(d,['Verify the active AWS identity with STS.','Review the IAM account summary, credential report, users, groups, roles, policies and MFA indicators where permitted.','Confirm S3 and CloudFront inventory after the previous hosting teardown.','Review recent CloudTrail event history where permitted.','Review current billing data where permitted.','Capture and redact evidence, record AccessDenied as a finding, and prepare remediation recommendations without applying them.'])
    d.add_heading('Exclusions and agreement controls',2); bullets(d,['No IAM, EC2, S3, CloudFront, VPC, database, logging, backup, or monitoring resource will be created.','No permissions, policies, credentials, MFA devices, budgets, trails, or account settings will be changed.','No access key, secret key, session token, password, MFA seed, full account number, or private email will enter the portfolio package.','This is training evidence and must not be represented as paid client or production experience.'])
    d.add_heading('Architecture and trust boundaries',1)
    d.add_paragraph('Operator with MFA -> AWS Console or CloudShell -> STS identity verification -> IAM, S3, CloudFront, CloudTrail and Billing read-only queries -> local redaction workspace -> independent AI review -> portfolio case study.')
    table(d,['Boundary','Allowed flow','Control'],[
        ['Operator device to AWS','HTTPS console and CloudShell session','MFA-protected operator; temporary authenticated session'],['AWS account services','Read-only List, Get, Describe and Lookup operations','No Create, Put, Update, Attach, Detach or Delete actions'],['AWS results to evidence','Screenshots and copied query output','Redact identifiers and exclude credentials'],['Evidence to external AI','Consolidated review package','Share only approved redacted content']],[1.8,2.4,2.7])
    d.add_heading('Project plan',1)
    table(d,['Phase','Activities','Output'],[
        ['1 Scope and access','Confirm read-only rule, identity, evidence naming and redaction standard','Approved scope record'],['2 Inventory','Run IAM, S3 and CloudFront checks','Inventory and screenshots'],['3 Validate','Review MFA, credentials, CloudTrail and billing','Test results and findings'],['4 Review','Export one-form AI package and record corrections','Independent review record'],['5 Handover','Finalize portfolio case study and lessons learned','Portfolio dossier']],[1.35,3.5,2.05])
    d.add_heading('Practical AWS console runbook',1)
    numbered(d,['Sign in as the MFA-protected operator, not the root user.','Open CloudShell and run aws sts get-caller-identity. Confirm the expected identity, then redact the account number in evidence.','Run aws iam get-account-summary. Record the root MFA indicator and account-level counts.','Run aws iam generate-credential-report, wait for completion, then run aws iam get-credential-report --query Content --output text. Do not publish the raw report.','Run aws iam list-users, aws iam list-groups, aws iam list-roles and aws iam list-policies --scope Local. Record only security-relevant summaries.','Run aws s3api list-buckets --query "Buckets[].Name". Expected post-teardown result is an empty list.','Run aws cloudfront list-distributions --query "DistributionList.Items[].{Id:Id,Domain:DomainName,Enabled:Enabled}". Expected post-teardown result is null or an empty list.','Open CloudTrail Event history or run aws cloudtrail lookup-events --max-results 20. Record successful activity or AccessDenied. Do not enable a trail.','Open Billing and Cost Management. Review current-month charges and credits for S3, CloudFront and EC2. Do not claim zero cost until the billing view supports it.','Complete the test table, redact screenshots, export the one-form AI verification package, paste the independent findings back into the application, correct the dossier and retest.'])
    d.add_heading('Command safety rule',2)
    d.add_paragraph('Before pressing Enter, confirm the command begins with aws and uses only get, list, describe or lookup semantics. Stop if a command contains create, put, update, attach, detach, enable, disable, delete, terminate or modify.')
    d.add_heading('Test and evidence matrix',1)
    rows=[
        ['T01','Caller identity','Expected operator returned','Pending','01-sts-caller-redacted.png'],['T02','Root MFA indicator','MFA enabled value recorded','Pending','02-iam-summary-redacted.png'],['T03','Credential hygiene','Credential report reviewed locally','Pending','03-credential-findings.png'],['T04','S3 inventory','No buckets after teardown','Pending','04-s3-inventory.png'],['T05','CloudFront inventory','No distributions after teardown','Pending','05-cloudfront-inventory.png'],['T06','Audit history','Events reviewed or AccessDenied recorded','Pending','06-cloudtrail-review.png'],['T07','Billing review','Current charges and credits recorded','Pending','07-billing-redacted.png'],['T08','No resource creation','Final inventory unchanged','Pending','08-final-attestation.png']]
    table(d,['ID','Test','Expected result','Actual result','Screenshot'],rows,[.55,1.35,2.25,1.05,1.75])
    d.add_heading('Screenshot standard',1)
    bullets(d,['Capture the whole relevant AWS panel so the service and result are visible.','Hide account number, email address, user IDs, ARNs where unnecessary, access keys, tokens, passwords and MFA details.','Use the exact filenames in the test matrix.','Do not mark a test Passed until the screenshot is reviewed for redaction and matches the expected result.','Store screenshots in the linked portfolio evidence workspace; use captions explaining what each image proves.'])
    d.add_heading('Findings and remediation record',1)
    table(d,['ID','Finding','Severity','Evidence','Recommendation','Status'],[
        ['F01','Permissions prevent one or more checks','To assess','AccessDenied screenshot','Treat as least-privilege evidence; request only the minimum read permission if essential','Open'],['F02','Long-lived credentials or inactive identities','To assess','Credential report finding','Disable or rotate only through a separately approved change','Open'],['F03','Residual AWS charge after teardown','To assess','Billing view','Identify service, region and usage date before remediation','Open'],['F04','Sensitive data appears in evidence','High','Evidence review','Reject the screenshot, redact it and recapture','Preventive']],[.4,1.25,.65,1.2,2.75,.65])
    d.add_heading('External AI verification form',1)
    d.add_paragraph('Reviewer instruction. Independently assess scope, architecture, command safety, IAM reasoning, cost claims, test coverage, evidence quality and portfolio wording. Treat every unreferenced claim as unverified. Never request secrets. Return PASS, PASS WITH CONDITIONS or FAIL, followed by blocking findings, non-blocking improvements, contradictions, exact corrections and retests.')
    table(d,['Reviewer field','Response'],[['Reviewer or tool',''],['Review date',''],['Verdict','NOT REVIEWED'],['Blocking findings',''],['Improvements',''],['Corrections applied',''],['Retest results',''],['Residual risks','']],[1.55,5.35])
    d.add_heading('Handover and portfolio case study',1)
    d.add_paragraph('Portfolio statement after evidence is complete. I performed a read-only AWS identity and account security assessment using STS, IAM, S3, CloudFront, CloudTrail and Billing views. I documented access boundaries, credential and MFA observations, post-teardown resource inventory, cost checks, redacted evidence, remediation recommendations and independent review results without creating or modifying AWS resources.')
    d.add_heading('Completion checklist',2); bullets(d,['All eight tests have actual results and reviewed screenshots.','Every finding references evidence and separates observation from recommendation.','The independent review response is saved in the application.','Corrections are retested and remaining risks are stated.','The portfolio contains no secrets or unsupported production claims.','Final AWS inventory confirms the assessment created no resources.'])
    path=OUT/'David Gaisey-Otoo AWS Identity and Account Security Assessment.docx'; d.save(path); return path

def build_challenges():
    d=Document(); setup(d,'AWS Assessment Challenges and Improvement Register','Separate working record for application and delivery improvements')
    d.add_heading('Purpose',1); d.add_paragraph('This register is separate from the public portfolio. It records limitations, defects, decisions and improvements discovered while AWS Career Launchpad Pro generated and validated the assessment workflow.')
    table(d,['ID','Challenge','Impact','Decision or correction','Timing','Status'],[
        ['C01','The original read-only brief generated IAM role creation steps and CloudFormation.','Could cause an unauthorized write.','Added aws-read-only detection, disabled IaC and replaced the runbook with query-only checks.','Before assessment','Fixed'],['C02','External review existed only inside the later delivery package.','The owner could not verify the whole solution in one place.','Added a consolidated editable AI verification package and saved reviewer-response field in Solution Studio.','Before assessment','Fixed'],['C03','The application cannot independently prove AWS results without an authenticated account session and evidence.','Generated output could be mistaken for completed work.','Require actual-result fields, screenshot references and post-check inventory before portfolio completion.','During assessment','Open control'],['C04','Some IAM, CloudTrail or Billing checks may return AccessDenied.','Coverage may be incomplete.','Record the denial as evidence; do not expand privileges automatically. Request minimum read access only if essential.','During assessment','Expected'],['C05','Billing data can arrive after resource deletion.','A same-day zero value may not be final.','Record the review date and recheck after billing data settles before claiming no residual charge.','After assessment','Pending'],['C06','Screenshots can expose account IDs, ARNs and email addresses.','Portfolio privacy and security risk.','Use the evidence naming and redaction gate before every upload or external review.','Every capture','Open control'],['C07','The app’s browser data is device-local unless synchronization is configured.','Saved reviews or evidence may not appear on another device.','Export the Markdown review package and retain approved documents in the categorized document vault.','Before handover','Improvement'],['C08','Email or SMS delivery has not been authorized or configured for this engagement.','Automatic sending could disclose unfinished material.','Keep communications as drafts until David explicitly approves the recipient and send action.','Handover','Open control']],[.5,1.65,1.35,2.55,.85,.75])
    d.add_heading('Decision gate',1); d.add_paragraph('Proceed with the read-only assessment only after the generated solution shows AWS read-only mode, one-click deployment disabled, no CloudFormation template, and a query-only runbook. Stop and log a new challenge if any screen proposes a write action or requests permanent credentials.')
    d.add_page_break(); d.add_heading('Improvement intake',1)
    table(d,['Date','Source','Finding','Priority','Owner','Target release','Retest evidence'],[['','','','','David Gaisey-Otoo','',''],['','','','','David Gaisey-Otoo','',''],['','','','','David Gaisey-Otoo','','']],[.65,1.0,2.2,.7,1.15,1.0,1.2])
    path=OUT/'David Gaisey-Otoo AWS Assessment Challenges and Improvements.docx'; d.save(path); return path

if __name__=='__main__':
    for p in (build_main(),build_challenges()): print(p.resolve())
